"""Genera una presentación explicativa editable a partir de SD-04.

No sustituye las tres presentaciones formales del Art. 45. La fuente de
contenido es latex_final/sd-04.tex y el informe explicativo asociado.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "05_Gestion/reportes/presentacion_explicativa_sd04.pptx"

NAVY = RGBColor(22, 42, 65)
BLUE = RGBColor(28, 104, 150)
TEAL = RGBColor(18, 127, 120)
PURPLE = RGBColor(109, 77, 143)
GOLD = RGBColor(176, 120, 27)
RED = RGBColor(174, 70, 67)
INK = RGBColor(35, 51, 67)
MUTED = RGBColor(92, 109, 124)
PALE = RGBColor(245, 248, 251)
PALE_BLUE = RGBColor(231, 242, 249)
PALE_TEAL = RGBColor(229, 244, 240)
PALE_PURPLE = RGBColor(240, 234, 248)
PALE_GOLD = RGBColor(251, 243, 225)
PALE_RED = RGBColor(250, 235, 234)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(209, 220, 228)
FONT = "DejaVu Sans"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def rect(slide, x, y, w, h, fill=WHITE, stroke=LINE, radius=True):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = stroke
    shp.line.width = Pt(1)
    return shp


def txt(slide, x, y, w, h, value, size=20, color=INK, bold=False,
        align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(.03)
    frame.margin_right = Inches(.03)
    frame.margin_top = Inches(.01)
    frame.margin_bottom = Inches(.01)
    frame.vertical_anchor = valign
    for idx, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(5)
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
    return box


def card(slide, x, y, w, h, title, body, accent=BLUE, tint=PALE_BLUE,
         body_size=18, title_size=20):
    rect(slide, x, y, w, h, tint, tint)
    rect(slide, x, y, .075, h, accent, accent, False)
    txt(slide, x+.25, y+.13, w-.48, .48, title, title_size, accent, True)
    txt(slide, x+.25, y+.69, w-.48, h-.79, body, body_size, INK,
        valign=MSO_ANCHOR.TOP)


def pill(slide, x, y, w, label, color=BLUE, tint=PALE_BLUE):
    rect(slide, x, y, w, .43, tint, tint)
    txt(slide, x+.10, y, w-.20, .43, label, 13, color, True, PP_ALIGN.CENTER)


def arrow(slide, x1, y1, x2, y2, color=BLUE, width=2, dash=None):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return shape


def slide(title, subtitle="", source="SD-04", notes=""):
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, WHITE, WHITE, False)
    rect(s, 0, 0, .16, 7.5, TEAL, TEAL, False)
    txt(s, .49, .34, 12.25, .48, title, 29, NAVY, True)
    if subtitle:
        txt(s, .50, .85, 12.25, .39, subtitle, 15, MUTED)
    rect(s, .50, 6.93, 12.26, .015, LINE, LINE, False)
    txt(s, .52, 6.99, 10.2, .26,
        f"Only Simple Solutions  ·  TFEP-01/2026  ·  Fuente: {source}",
        10, MUTED)
    txt(s, 12.0, 6.98, .7, .28, f"{len(prs.slides):02}",
        11, MUTED, True, PP_ALIGN.RIGHT)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def lead(s, value):
    txt(s, .55, 1.32, 12.12, .59, value, 20, INK, True)


# 1 — Portada
s = prs.slides.add_slide(blank)
rect(s, 0, 0, 13.333, 7.5, NAVY, NAVY, False)
rect(s, 0, 0, .20, 7.5, TEAL, TEAL, False)
txt(s, .8, .68, 10.9, .44, "ONLY SIMPLE SOLUTIONS  ·  LICITACIÓN TFEP-01/2026",
    17, RGBColor(169, 217, 221), True)
txt(s, .8, 1.58, 11.5, 1.6, "Subdocumento 4\nArquitectura lógica y física", 35, WHITE, True)
txt(s, .82, 3.78, 10.8, .76,
    "Qué está diseñado, por qué se eligió y qué falta verificar", 23, WHITE)
pill(s, .83, 5.05, 3.58, "Diseño lógico definido", TEAL, PALE_TEAL)
pill(s, 4.57, 5.05, 4.13, "Diseño físico preliminar", GOLD, PALE_GOLD)
txt(s, .84, 6.69, 11.7, .34, "Revisión de trabajo · 8 de octubre de 2026",
    15, RGBColor(189, 203, 217))
s.notes_slide.notes_text_frame.text = (
    "Esta presentación explica la versión actual del SD-04. No es una selección final "
    "de proveedor, productos o topología física. El informe explicativo y el LaTeX "
    "actual contienen el detalle completo."
)

# 2 — Problema
s = slide("El problema: cuatro promesas que deben coincidir",
          "Ancoa es una tienda y una filial emisora de crédito",
          "Caso 09; SD-04 §4",
          "El Caso describe plataformas conectadas mediante catorce interfaces y estados "
          "inconsistentes. SD-04 organiza responsables y recorridos para que precio, "
          "existencia, entrega y crédito sean verificables sin mezclar datos.")
for x, n, label, col, tint in [
    (.65, "01", "El producto\nexiste", TEAL, PALE_TEAL),
    (3.8, "02", "El precio\nes el exhibido", BLUE, PALE_BLUE),
    (6.95, "03", "La entrega\nllega a tiempo", GOLD, PALE_GOLD),
    (10.1, "04", "El crédito\nes el informado", PURPLE, PALE_PURPLE),
]:
    rect(s, x, 2.02, 2.55, 2.35, tint, tint)
    txt(s, x+.24, 2.20, .75, .48, n, 22, col, True)
    txt(s, x+.23, 2.85, 2.12, 1.17, label, 21, NAVY, True)
card(s, .67, 4.76, 5.90, 1.36, "Situación actual",
     "22 tiendas · 14 interfaces · dos negocios", RED, PALE_RED, 18, 19)
card(s, 6.77, 4.76, 5.90, 1.36, "Papel del SD-04",
     "Asignar responsable, flujo, dato y ubicación", BLUE, PALE_BLUE, 18, 19)

# 3 — Alcance
s = slide("Qué cambia y qué permanece",
          "La arquitectura respeta el alcance ya establecido en SD-03",
          "SD-04 tabla 4.4; SD-03 §3.2",
          "ERP/DTE sigue emitiendo documentos tributarios; WMS conserva ejecución "
          "física y marketplace su plataforma. Ecommerce y fidelización se integran "
          "inicialmente; su destino se decide con pruebas. Núcleo Retail, crédito y POS "
          "se reemplazan por olas conciliadas. El motor de precios no existe hoy.")
card(s, .63, 1.55, 3.88, 3.30, "CONSERVAR E INTEGRAR",
     "ERP / DTE\nWMS principal\nMarketplace 2022", TEAL, PALE_TEAL)
card(s, 4.72, 1.55, 3.88, 3.30, "INTEGRAR Y EVALUAR",
     "Comercio electrónico\nFidelización\nDecisión tras pruebas", BLUE, PALE_BLUE)
card(s, 8.81, 1.55, 3.88, 3.30, "REEMPLAZAR POR OLAS",
     "Núcleo Retail 2009\nCrédito 2011\nPOS de tienda 2014", PURPLE, PALE_PURPLE)
rect(s, .65, 5.18, 12.03, .93, PALE_GOLD, PALE_GOLD)
txt(s, .88, 5.32, 11.6, .59,
    "Nueva capacidad: motor de precios. El Caso lo menciona, pero no existe como plataforma vigente.",
    18, INK)

# 4 — arquitectura
s = slide("Del actor al dato: el recorrido lógico",
          "Las flechas muestran funciones, no enlaces físicos ya acreditados",
          "SD-04 figuras 4.1–4.2",
          "La entrada se controla por gateway. Los servicios de Retail y Emisor "
          "conservan autoridad distinta. Adaptadores conectan plataformas vigentes; "
          "broker distribuye hechos confirmados y no confirma reservas ni crédito. "
          "La tienda mantiene diario local cuando cae la red.")
stages = [
    (.63, "ACTORES", "Cliente · cajero\nanalista · sistemas", BLUE, PALE_BLUE),
    (3.18, "CANALES", "POS · web\nconsolas · API", TEAL, PALE_TEAL),
    (5.73, "GATEWAY", "Identidad · cuotas\ncontratos · trazas", BLUE, PALE_BLUE),
    (8.28, "SERVICIOS", "Retail · Emisor\nFrontera X-01", PURPLE, PALE_PURPLE),
    (10.83, "DATOS", "Autoridades\nseparadas", GOLD, PALE_GOLD),
]
for x, title, body, col, tint in stages:
    card(s, x, 2.05, 1.93, 1.96, title, body, col, tint, 14, 16)
for x in [2.65, 5.20, 7.75, 10.30]:
    arrow(s, x, 3.03, x+.45, 3.03, BLUE, 2)
card(s, 2.14, 4.76, 4.2, 1.21, "TIENDA LOCAL",
     "Diario y sincronizador para cortes", TEAL, PALE_TEAL, 16, 18)
card(s, 7.02, 4.76, 4.2, 1.21, "INTEGRACIÓN",
     "Adaptadores + broker candidato", BLUE, PALE_BLUE, 16, 18)

# 5 — vistas
s = slide("Cinco vistas explican el mismo sistema",
          "Exigencia de RT-02.03; cada vista responde una pregunta diferente",
          "Bases Transversales RT-02.03; SD-04 §4.1",
          "RT-02.03 exige vistas lógica, de procesos, de despliegue, de datos "
          "y de seguridad. SD-04 mantiene los mismos códigos para que una "
          "responsabilidad no cambie entre diagramas.")
views = [
    (.65, 1.66, "LÓGICA", "¿Quién hace qué?", BLUE, PALE_BLUE),
    (4.74, 1.66, "PROCESOS", "¿En qué orden ocurre?", TEAL, PALE_TEAL),
    (8.83, 1.66, "DESPLIEGUE", "¿Dónde se ejecuta?", GOLD, PALE_GOLD),
    (2.69, 4.06, "DATOS", "¿Quién escribe y lee?", PURPLE, PALE_PURPLE),
    (6.78, 4.06, "SEGURIDAD", "¿Quién puede acceder?", RED, PALE_RED),
]
for x, y, title, body, col, tint in views:
    card(s, x, y, 3.86, 1.76, title, body, col, tint, 18, 20)

# 6 — capas
s = slide("Ocho capas, con responsabilidades distintas",
          "Seguridad y observabilidad acompañan a todas las demás",
          "Bases Transversales RT-02.01; SD-04 figura 4.2",
          "Las ocho responsabilidades vienen de RT-02.01. No son ocho servidores "
          "ni pasos obligatorios por los que pase cada transacción. Los actores "
          "de sistemas entran por contratos o adaptadores, sin acceso directo "
          "a bases de datos nuevas.")
layers = [
    ("1", "Presentación", "Pantallas y POS"),
    ("2", "Borde", "Exposición protegida"),
    ("3", "Gateway", "API autorizadas"),
    ("4", "Servicios", "Reglas Retail / Emisor"),
    ("5", "Integración", "Contratos y eventos"),
    ("6", "Datos", "Registros por autoridad"),
]
for i, (num, name, body) in enumerate(layers):
    x = .69 + (i%3)*4.20
    y = 1.54 + (i//3)*1.77
    rect(s, x, y, 3.92, 1.48, PALE_BLUE if i%2==0 else PALE_TEAL,
         PALE_BLUE if i%2==0 else PALE_TEAL)
    txt(s, x+.17, y+.19, .46, .43, num, 20, BLUE, True)
    txt(s, x+.72, y+.14, 2.95, .48, name, 20, NAVY, True)
    txt(s, x+.72, y+.72, 3.0, .42, body, 17, INK)
pill(s, 1.50, 5.49, 4.54, "7 · Seguridad transversal", RED, PALE_RED)
pill(s, 7.30, 5.49, 4.54, "8 · Observabilidad transversal", PURPLE, PALE_PURPLE)

# 7 — servicios
s = slide("Trece responsabilidades; despliegue selectivo",
          "Una responsabilidad de negocio no equivale siempre a un microservicio",
          "SD-04 tablas 4.2–4.3",
          "Nueve responsabilidades son Retail, tres Emisor y una gobierna los "
          "cruces. Ventas, existencias y crédito requieren autoridad y recuperación "
          "propias. Comisiones y gobierno marketplace pueden comenzar como módulos "
          "extraíbles. ERP sigue calculando remuneraciones y emitiendo DTE; WMS "
          "sigue ejecutando el almacén.")
card(s, .63, 1.61, 5.10, 3.60, "9 · RETAIL",
     "Oferta · abastecimiento · existencias\nPedidos · ventas · comisiones\nMarketplace · posventa · clientes",
     TEAL, PALE_TEAL, 18, 23)
card(s, 5.95, 1.61, 3.18, 3.60, "3 · EMISOR",
     "Originación\nCartera\nEvidencia financiera",
     PURPLE, PALE_PURPLE, 18, 23)
card(s, 9.36, 1.61, 3.33, 3.60, "1 · FRONTERA",
     "X-01 gobierna\npermisos y bitácora\nde cruces",
     GOLD, PALE_GOLD, 18, 23)
rect(s, .65, 5.51, 12.03, .61, PALE_BLUE, PALE_BLUE)
txt(s, .89, 5.57, 11.55, .47,
    "Criterio: separar cuando autoridad, carga o recuperación lo justifiquen.",
    18, NAVY, True)

# 8 — actores
s = slide("La cobertura parte de 31 actores UAW",
          "18 roles humanos + 13 sistemas o disparadores; AS-06 no existe en el catálogo",
          "SD-04 tabla 4.1; catálogo UAW",
          "El catálogo completo está trazado en el informe de especificación de "
          "diagramas. AH-17 usa API por decisión UAW; AS-11 es remuneraciones "
          "dentro del ERP y AS-14 solo dispara tareas. Actor no significa servidor.")
examples = [
    ("Cliente / cajero", "POS o web → ventas y existencias", TEAL, PALE_TEAL),
    ("Personal logístico", "Terminal → abastecimiento y WMS", BLUE, PALE_BLUE),
    ("Titular / ejecutivo", "Rol financiero → servicios Emisor", PURPLE, PALE_PURPLE),
    ("Vendedor externo", "API → gobierno marketplace", GOLD, PALE_GOLD),
    ("Analista BI", "Autoservicio → ámbito permitido", RED, PALE_RED),
]
for i, (title, body, col, tint) in enumerate(examples):
    y = 1.52 + i*1.00
    rect(s, .76, y, 11.82, .82, tint, tint)
    rect(s, .76, y, .075, .82, col, col, False)
    txt(s, 1.02, y+.13, 3.36, .54, title, 18, col, True)
    txt(s, 4.46, y+.13, 7.70, .54, body, 18, INK)

# 9 — venta conectada
s = slide("Venta conectada: confirmar antes de informar",
          "La reserva y el crédito necesitan respuesta antes de prometer o cobrar",
          "SD-04 §4.1, recorrido de existencia y venta",
          "R:M-01 publica precio vigente; R:M-03 confirma reserva. R:V-02 "
          "registra venta y coordina pago. ERP/DTE es el único emisor tributario. "
          "El hecho confirmado se distribuye después. El plazo de disponibilidad "
          "publicada es 30 segundos según el diseño del documento.")
steps = [
    ("1", "Cajero", "inicia venta"),
    ("2", "Oferta", "precio vigente"),
    ("3", "Existencias", "reserva confirmada"),
    ("4", "Ventas", "registra y cobra"),
    ("5", "ERP / DTE", "documento fiscal"),
    ("6", "Evento", "actualiza vistas"),
]
for i, (n, title, body) in enumerate(steps):
    x=.61+(i%3)*4.21
    y=1.67+(i//3)*2.07
    col=TEAL if i<3 else BLUE
    tint=PALE_TEAL if i<3 else PALE_BLUE
    rect(s,x,y,3.91,1.55,tint,tint)
    txt(s,x+.18,y+.17,.46,.40,n,19,col,True)
    txt(s,x+.72,y+.12,2.99,.48,title,20,NAVY,True)
    txt(s,x+.72,y+.68,2.95,.46,body,17,INK)
txt(s,.81,6.05,11.8,.44,
    "Regla: sin reserva confirmada no se promete disponibilidad ni se cobra la venta conectada.",
    17,NAVY,True)

# 10 — sin enlace
s = slide("Si cae la red, la tienda conserva su operación",
          "24 horas exigidas por RT-03.10; regreso con conciliación",
          "SD-04 figura 4.3; Bases Transversales RT-03.10",
          "El Caso menciona ocho horas de tienda, pero el requisito transversal "
          "exige veinticuatro. El POS mantiene oferta y existencias locales, escribe "
          "una venta pendiente en diario durable y sincroniza con clave idempotente "
          "al reconectar. Ventas confirma y publica después. DTE y cobro sin red "
          "requieren modalidad fiscal y operativa aprobada. No se originan tarjetas "
          "ni se amplían cupos sin conexión.")
flow=[
    ("1 · CORTE", "POS consulta datos locales", TEAL, PALE_TEAL),
    ("2 · REGISTRO", "Diario guarda venta pendiente", BLUE, PALE_BLUE),
    ("3 · RETORNO", "Sincronizador reintenta con ID", GOLD, PALE_GOLD),
    ("4 · CONCILIACIÓN", "Ventas confirma; luego publica", PURPLE, PALE_PURPLE),
]
for i,(title,body,col,tint) in enumerate(flow):
    x=.69+i*3.18
    card(s,x,2.03,2.91,2.74,title,body,col,tint,18,18)
    if i<3: arrow(s,x+2.92,3.40,x+3.16,3.40,col,2)
rect(s,.72,5.24,11.94,.80,PALE_RED,PALE_RED)
txt(s,.94,5.34,11.50,.56,
    "Abierto: modalidad de cobro y DTE durante el corte; crédito nuevo y aumento de cupo no se habilitan.",
    17,RED,True)

# 11 — sincronía
s = slide("Gateway, broker y diario cumplen tareas distintas",
          "Kafka es preferencia del equipo; el broker aún no está seleccionado",
          "SD-04 §4.1 y ADR-03/06",
          "El gateway gobierna solicitudes API, el broker difunde hechos ya "
          "confirmados y el diario conserva operaciones antes de que exista red. "
          "Kafka, Event Hubs y Service Bus se compararán por relectura, semántica "
          "de entrega, errores y recuperación. Ninguno sustituye respuesta síncrona "
          "para reservas o autorización de crédito.")
card(s,.68,1.70,3.83,3.29,"GATEWAY · AHORA",
     "Recibe API\nVerifica identidad y contrato\nDevuelve respuesta",BLUE,PALE_BLUE,18,20)
card(s,4.75,1.70,3.83,3.29,"BROKER · DESPUÉS",
     "Reparte hechos confirmados\nReintentos y consumidores\nProducto por decidir",PURPLE,PALE_PURPLE,18,20)
card(s,8.82,1.70,3.83,3.29,"DIARIO · SIN RED",
     "Guarda venta pendiente\nEs local y durable\nConcilia al reconectar",TEAL,PALE_TEAL,18,20)
txt(s,.80,5.42,11.73,.64,
    "El broker de nube no puede conservar un evento que nunca recibió de la tienda.",
    19,NAVY,True,PP_ALIGN.CENTER)

# 12 — frontera
s = slide("Retail y Emisor: datos separados",
          "X-01 documenta y controla cada cruce permitido",
          "SD-04 figura 4.5; ADR-02",
          "El comercio conserva identidad y preferencias; el Emisor conserva "
          "cupo, saldo, mora y evidencia crediticia. X-01 versiona finalidad, "
          "fundamento, campos y responsables. La política se aplica en ambos "
          "extremos. Una compra con tarjeta usa referencia, importe, medio y "
          "decisión mínima, sujeta a aprobación. Marketing no recibe mora ni pagos.")
card(s,.70,1.68,4.27,3.10,"RETAIL",
     "Venta · precio · pedido\nIdentidad y preferencias\nBI comercial",TEAL,PALE_TEAL,18,22)
card(s,8.36,1.68,4.27,3.10,"EMISOR",
     "Cupo · saldo · mora\nEvidencia financiera\nBI financiero",PURPLE,PALE_PURPLE,18,22)
rect(s,5.22,2.17,2.89,2.10,PALE_GOLD,PALE_GOLD)
txt(s,5.42,2.33,2.51,.44,"X-01",23,GOLD,True,PP_ALIGN.CENTER)
txt(s,5.42,2.96,2.51,.89,"Finalidad · mínimo\nbitácora",17,INK,False,PP_ALIGN.CENTER)
arrow(s,4.98,3.18,5.21,3.18,GOLD,2)
arrow(s,8.10,3.18,8.34,3.18,GOLD,2)
rect(s,.72,5.15,11.91,.79,PALE_RED,PALE_RED)
txt(s,.94,5.27,11.49,.55,
    "Acceso denegado: Marketing no consulta saldos, mora ni historial de pago.",
    18,RED,True,PP_ALIGN.CENTER)

# 13 — datos e IA
s = slide("Analítica útil sin abrir la frontera de datos",
          "BI segregado; modelos de inventario solo recomiendan",
          "SD-04 §§4.1–4.1.1",
          "Retail y Emisor cuentan con lagos, permisos y tableros separados. "
          "Una vista cruzada está deshabilitada salvo ficha y aprobación. Los "
          "modelos de conteo predictivo y clasificación de merma usan datos "
          "autorizados de Retail y proponen acciones; una persona revisa y "
          "R:M-03 registra el ajuste. Validar calidad y beneficio antes de desplegar.")
card(s,.70,1.56,5.79,2.05,"LAGO + BI RETAIL",
     "Precio, venta, conteos y existencias\nAcceso por rol comercial",TEAL,PALE_TEAL,18,20)
card(s,6.83,1.56,5.79,2.05,"LAGO + BI EMISOR",
     "Cartera y operación crediticia\nAcceso por rol financiero",PURPLE,PALE_PURPLE,18,20)
card(s,.70,3.96,5.79,2.05,"ML · CONTEO PREDICTIVO",
     "Sugiere dónde contar primero\nNo modifica inventario",BLUE,PALE_BLUE,18,20)
card(s,6.83,3.96,5.79,2.05,"ML · CLASIFICACIÓN DE MERMA",
     "Sugiere causa de la diferencia\nRevisión humana antes del ajuste",GOLD,PALE_GOLD,18,20)

# 14 — sitios
s = slide("Sitios distintos: tienda, distribución, datos y nube",
          "La ubicación física de cada plataforma conservada sigue por verificar",
          "SD-04 figura 4.4 y tabla 4.8",
          "SIT-01 agrupa 22 tiendas. CD-01 y CD-02 son centros de distribución; "
          "la operación WMS en el principal no acredita que su servidor esté "
          "allí. DC-01 es el centro de datos de 140 m² y DC-02 la sala de respaldo "
          "en la misma comuna. CLD-01 es nube primaria candidata en Santiago y "
          "CLD-02 región de recuperación candidata en São Paulo. La filial es "
          "una entidad, no un lugar físico.")
sites=[
    (.65,1.54,"SIT-01","22 tiendas\nPOS + diario",TEAL,PALE_TEAL),
    (4.72,1.54,"CD-01 / 02","Distribución principal\ny Concepción",BLUE,PALE_BLUE),
    (8.79,1.54,"DC-01 / 02","Centro de datos 140 m²\ny sala de respaldo",GOLD,PALE_GOLD),
    (2.69,4.04,"CLD-01 / 02","Nube Santiago\ny recuperación candidata",PURPLE,PALE_PURPLE),
    (6.77,4.04,"EXT-01","Pagos, transportistas,\nproveedores y otros",RED,PALE_RED),
]
for x1, x2 in [(2.57, 4.61), (6.64, 4.61), (10.71, 4.61)]:
    arrow(s, x1, 3.35, x2, 4.01, MUTED, 1.5, dash=True)
arrow(s, 6.54, 4.93, 6.73, 4.93, MUTED, 1.5, dash=True)
for x,y,title,body,col,tint in sites:
    card(s,x,y,3.84,1.78,title,body,col,tint,17,19)
txt(s,.81,6.14,11.65,.36,
    "Rutas y enlaces: propuestos en el diagrama; deben confirmarse con el inventario WAN.",
    15,MUTED)

# 15 — nube
s = slide("Azure es la referencia; Google sigue en comparación",
          "La nube y el broker no están adjudicados",
          "SD-04 §4.1.1 y §§4.2–4.3",
          "Azure es referencia por la alianza planteada en SD-01 y un conjunto "
          "candidato de servicios. La certificación de socio aún debe acreditarse. "
          "Google Cloud se compara, entre otros motivos, por Kafka administrado "
          "en Santiago. En Azure Chile Central no está acreditado Kafka "
          "administrado: operarlo en AKS implica carga de operación. Event Hubs "
          "usa protocolo Kafka pero no es Apache Kafka. Se verifican productos "
          "por región, residencia, recuperación, conectividad y soporte.")
card(s,.71,1.68,5.72,3.25,"AZURE · REFERENCIA",
     "Chile Central primaria\nBrazil South recuperación candidata\nAlianza y productos por acreditar",BLUE,PALE_BLUE,18,21)
card(s,6.88,1.68,5.72,3.25,"GOOGLE · ALTERNATIVA",
     "Santiago y São Paulo por comparar\nKafka administrado en Santiago\nComprobar el conjunto completo",TEAL,PALE_TEAL,18,21)
rect(s,.73,5.25,11.86,.91,PALE_GOLD,PALE_GOLD)
txt(s,.95,5.38,11.43,.62,
    "Cierre: pruebas por producto y región + residencia financiera + recuperación + operación híbrida.",
    18,NAVY,True,PP_ALIGN.CENTER)

# 16 — tecnología
s = slide("Tecnologías candidatas, no compras cerradas",
          "T-11 precisará versiones, cantidades, licencias y soporte",
          "SD-04 tabla 4.6; Formulario T-11",
          "Java 25 LTS y Spring Boot 4 son base candidata para API transaccionales; "
          "React/TypeScript para nuevas vistas; AKS para cargas propias; API "
          "Management para gateway; PostgreSQL y Redis reconstruible; Data Lake "
          "Storage y Power BI por ámbito; Entra y Key Vault sujetos a federación; "
          "OpenTelemetry, Azure Monitor y Terraform para operación. Todos tienen "
          "verificaciones de región, competencias, escala o compatibilidad.")
tech=[
    ("Aplicación", "Java 25 · Spring Boot 4 · React", TEAL, PALE_TEAL),
    ("Cómputo / API", "AKS · API Management", BLUE, PALE_BLUE),
    ("Mensajería", "Kafka / Event Hubs / Service Bus", PURPLE, PALE_PURPLE),
    ("Datos / BI", "PostgreSQL · Redis · Lake · Power BI", GOLD, PALE_GOLD),
    ("Identidad", "Federación · Entra ID · Key Vault", RED, PALE_RED),
    ("Operación", "OpenTelemetry · Monitor · Terraform", TEAL, PALE_TEAL),
]
for i,(title,body,col,tint) in enumerate(tech):
    x=.70+(i%2)*6.17
    y=1.45+(i//2)*1.58
    card(s,x,y,5.75,1.36,title,body,col,tint,16,19)
txt(s,.85,6.45,11.7,.29,
    "Pruebas pendientes: disponibilidad regional, compatibilidad, continuidad, competencias y seguridad.",
    13,MUTED)

# 17 — ADR
s = slide("Las seis decisiones y cómo se comprobarán",
          "Una decisión de arquitectura debe poder revisarse con evidencia",
          "SD-04 tabla 4.5",
          "ADR-01: olas y conciliación. ADR-02: datos separados y denegaciones. "
          "ADR-03: solicitudes síncronas para bloqueos, eventos para hechos. "
          "ADR-04: POS local probado 24 horas. ADR-05: microservicios selectivos "
          "según carga y autoridad. ADR-06: comparar broker y verificar gateway. "
          "Las decisiones aún no probadas no deben presentarse como definitivas.")
adrs=[
    ("01 · Migración por olas", "Mapa de interfaces y reversión", TEAL, PALE_TEAL),
    ("02 · Datos separados", "Permisos y denegación auditada", PURPLE, PALE_PURPLE),
    ("03 · Contratos mixtos", "Latencia, duplicados y reintento", BLUE, PALE_BLUE),
    ("04 · POS local", "Piloto 24 h y conciliación", TEAL, PALE_TEAL),
    ("05 · Partición selectiva", "Carga y falla aislada", GOLD, PALE_GOLD),
    ("06 · Broker + gateway", "Comparar semántica y recuperación", BLUE, PALE_BLUE),
]
for i,(title,body,col,tint) in enumerate(adrs):
    x=.70+(i%2)*6.17
    y=1.45+(i//2)*1.58
    card(s,x,y,5.75,1.36,title,body,col,tint,16,19)

# 18 — pendientes
s = slide("Lo que falta para cerrar la arquitectura física",
          "Un diagrama lógico no acredita dónde está instalado cada sistema",
          "SD-04 §§4.1–4.3; tabla 4.8",
          "Falta inventariar catorce interfaces, alojamiento de plataformas, "
          "WAN/enlaces, equipos y periféricos, informe de brechas DC 2024, "
          "volumen, retención, política comercial y fiscal de contingencia, "
          "cruces jurídicos, residencia financiera y pruebas regionales. Cada "
          "arista física final tendrá origen, destino, contrato, red, protección, "
          "volumen, latencia, responsable y comportamiento ante falla.")
pending=[
    ("01", "Interfaces y alojamiento"),
    ("02", "WAN, tiendas y periféricos"),
    ("03", "Brechas de DC principal y respaldo"),
    ("04", "Volumetría y capacidad"),
    ("05", "Cruces legales y residencia de datos"),
    ("06", "Pago, DTE y precio durante contingencia"),
]
for i,(n,body) in enumerate(pending):
    x=.70+(i%2)*6.17
    y=1.45+(i//2)*1.58
    rect(s,x,y,5.75,1.36,PALE_GOLD,PALE_GOLD)
    txt(s,x+.20,y+.16,.54,.42,n,19,GOLD,True)
    txt(s,x+.85,y+.13,4.66,.99,body,18,INK,True)
txt(s,.82,6.46,11.70,.30,
    "El siguiente paso es levantar evidencia y trasladarla a T-11 y al diagrama físico verificable.",
    13,MUTED)

# 19 — cierre
s = slide("Estado actual del subdocumento 4",
          "La propuesta ya puede explicarse; su topología final requiere comprobaciones",
          "SD-04 completo; informe explicativo",
          "Las cinco vistas exigidas, los 31 actores y los ocho emplazamientos "
          "están declarados. Quedan definidos los flujos de venta, corte, retorno "
          "y frontera. El producto del broker, el proveedor de nube y las rutas "
          "reales no están cerrados. La arquitectura física necesita inventario, "
          "pruebas y aprobación jurídica.")
card(s,.72,1.73,5.78,3.11,"YA DEFINIDO",
     "5 vistas · 31 actores\n13 responsabilidades · 8 sitios\nFlujos y controles de frontera",
     TEAL,PALE_TEAL,18,21)
card(s,6.84,1.73,5.78,3.11,"POR VALIDAR",
     "Proveedor y broker\nRutas, capacidad y productos\nContingencia fiscal y residencia",
     GOLD,PALE_GOLD,18,21)
rect(s,.74,5.21,11.84,.86,PALE_BLUE,PALE_BLUE)
txt(s,.96,5.35,11.39,.56,
    "Idea principal: el diseño separa responsabilidades; la evidencia cerrará el despliegue.",
    19,NAVY,True,PP_ALIGN.CENTER)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"{OUT} ({len(prs.slides)} diapositivas)")
